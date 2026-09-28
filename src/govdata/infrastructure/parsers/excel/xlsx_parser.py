"""Adaptador: XlsxParser -- implementa `DocumentParserPort` para XLSX.

`read_only=True` de openpyxl evita cargar el workbook completo en
memoria: itera filas bajo demanda directamente del XML comprimido.

Con `reglas_normalizacion` inyectadas elige la hoja y la fila de
encabezado reales (`detector_encabezado.py`): los Excel publicos suelen
traer titulos institucionales, logos o filas vacias arriba del
encabezado, celdas combinadas y varias hojas. En modo read_only las
celdas combinadas no-ancla llegan como `None` (celda vacia), lo que no
afecta la deteccion. Sin reglas se conserva el comportamiento historico
(hoja activa, primera fila).
"""

from __future__ import annotations

import itertools

import openpyxl

from govdata.application.dto.resultado_parseo_dto import EstadoParseo, ResultadoParseoDTO
from govdata.application.dto.tabla_extraida_dto import TablaExtraidaDTO
from govdata.application.ports.document_parser_port import DocumentParserPort
from govdata.config.logging_config import get_logger
from govdata.domain.entities.documento import TipoArchivo
from govdata.domain.rules.detector_encabezado import (
    FILAS_BUSQUEDA_ENCABEZADO,
    encabezados_unicos,
    limpiar_celda_encabezado,
    seleccionar_hoja_y_encabezado,
)
from govdata.domain.rules.reglas_configurables import ReglasNormalizacion

logger = get_logger(__name__)


class XlsxParser(DocumentParserPort):
    def __init__(self, reglas_normalizacion: ReglasNormalizacion | None = None) -> None:
        self._reglas_normalizacion = reglas_normalizacion

    def soporta(self, tipo_archivo: TipoArchivo) -> bool:
        return tipo_archivo == TipoArchivo.XLSX

    def extraer(self, ruta_archivo: str) -> ResultadoParseoDTO:
        try:
            libro = openpyxl.load_workbook(ruta_archivo, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - cualquier fallo de apertura es "corrupto"
            logger.error("xlsx_ilegible", ruta=ruta_archivo, error=str(exc))
            return ResultadoParseoDTO(estado=EstadoParseo.ERROR, mensaje=f"Excel (.xlsx) corrupto o ilegible: {exc}")

        try:
            hoja_activa = libro.active
            if hoja_activa is None:
                return ResultadoParseoDTO(
                    estado=EstadoParseo.EXITO,
                    tabla=TablaExtraidaDTO(encabezados=(), filas=[]),
                    mensaje="El libro no tiene ninguna hoja activa",
                )

            # La hoja activa va primera: ante empate, gana (compatibilidad).
            hojas = [hoja_activa, *(h for h in libro.worksheets if h is not hoja_activa)]
            indice_hoja, indice_fila = 0, 0
            if self._reglas_normalizacion is not None:
                muestras = [
                    list(itertools.islice(h.iter_rows(values_only=True), FILAS_BUSQUEDA_ENCABEZADO)) for h in hojas
                ]
                deteccion = seleccionar_hoja_y_encabezado(muestras, self._reglas_normalizacion)
                indice_hoja, indice_fila = deteccion.indice_hoja, deteccion.indice_fila
            hoja = hojas[indice_hoja]

            filas_iter = hoja.iter_rows(values_only=True)
            encabezados_raw = next(itertools.islice(filas_iter, indice_fila, None), None)
            if encabezados_raw is None:
                return ResultadoParseoDTO(
                    estado=EstadoParseo.EXITO,
                    tabla=TablaExtraidaDTO(encabezados=(), filas=[]),
                    mensaje="Excel vacio",
                )

            encabezados = encabezados_unicos([limpiar_celda_encabezado(celda) for celda in encabezados_raw])
            filas: list[dict[str, str]] = []
            for fila_raw in filas_iter:
                if fila_raw is None or all(celda is None or not str(celda).strip() for celda in fila_raw):
                    continue
                valores = [str(celda).strip() if celda is not None else "" for celda in fila_raw]
                filas.append(dict(zip(encabezados, valores, strict=False)))

            logger.info(
                "xlsx_extraido",
                ruta=ruta_archivo,
                filas=len(filas),
                hoja=hoja.title,
                fila_encabezado=indice_fila + 1,
            )
            return ResultadoParseoDTO(estado=EstadoParseo.EXITO, tabla=TablaExtraidaDTO(encabezados=encabezados, filas=filas))
        finally:
            libro.close()
