"""Genera las fixtures del Modulo 4 (Motor de Extraccion). Ejecutar una
sola vez (o tras cambiar el contenido esperado):

    python tests/fixtures/extraccion/_generar.py

No se ejecuta como parte de la suite -- los archivos generados quedan
versionados junto al codigo.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

FIXTURES = Path(__file__).resolve().parent

_FILAS_DIRECTORIO = [
    ["Nombre Completo", "Cargo", "Dependencia", "Correo", "Telefono"],
    ["Juan Carlos Perez Gomez", "Director Juridico", "Procuraduria", "jperez@entidad.gob.ec", "022345678"],
    ["Maria Elena Torres", "Contadora", "Financiero", "mtorres@entidad.gob.ec", "022345679"],
    ["Ana Lucia Ramos Vega", "Asesora Juridica", "Legal", "aramos@entidad.gob.ec", "022345680"],
    ["Pedro Jose Salas", "Director Financiero", "Financiero", "psalas@entidad.gob.ec", ""],
    ["Luis Fernando Diaz", "Chofer", "Transporte", "", ""],
]


def _generar_pdf_con_tabla() -> None:
    doc = SimpleDocTemplate(str(FIXTURES / "directorio_con_tabla.pdf"), pagesize=letter)
    tabla = Table(_FILAS_DIRECTORIO)
    tabla.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
    doc.build([tabla])


def _generar_pdf_con_texto() -> None:
    ruta = FIXTURES / "directorio_con_texto.pdf"
    c = canvas.Canvas(str(ruta))
    c.setFont("Courier", 9)
    anchos = [26, 20, 16, 26, 12]
    y = 750
    for fila in _FILAS_DIRECTORIO:
        linea = "".join(valor.ljust(ancho) for valor, ancho in zip(fila, anchos, strict=False))
        c.drawString(40, y, linea)
        y -= 18
    c.save()


def _generar_pdf_escaneado() -> None:
    ruta = FIXTURES / "directorio_escaneado.pdf"
    c = canvas.Canvas(str(ruta))
    c.rect(50, 50, 400, 700)  # solo grafico vectorial, cero texto
    c.save()


def _generar_pdf_vacio() -> None:
    ruta = FIXTURES / "pdf_vacio.pdf"
    c = canvas.Canvas(str(ruta))
    c.showPage()
    c.save()


def _generar_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Directorio"
    ws.append(["Servidor", "Puesto", "Unidad", "Correo Institucional", "Celular"])
    for fila in _FILAS_DIRECTORIO[1:]:
        ws.append(fila)
    ws.append([None, None, None, None, None])  # fila vacia -- debe filtrarse
    wb.save(FIXTURES / "directorio.xlsx")


def _generar_xlsx_columnas_desconocidas() -> None:
    """Columnas que NO coinciden con ningun sinonimo -- caso de error
    'columnas inexistentes' (todas las filas deben descartarse, sin
    tumbar el pipeline)."""
    wb = Workbook()
    ws = wb.active
    ws.append(["ColumnaX", "ColumnaY", "ColumnaZ"])
    ws.append(["dato1", "dato2", "dato3"])
    wb.save(FIXTURES / "columnas_desconocidas.xlsx")


def _generar_xlsx_vacio() -> None:
    wb = Workbook()
    wb.save(FIXTURES / "excel_vacio.xlsx")


def _generar_csv() -> None:
    ruta = FIXTURES / "directorio.csv"
    with ruta.open("w", newline="", encoding="utf-8") as archivo:
        escritor = csv.writer(archivo, delimiter=";")
        escritor.writerow(["Apellidos y Nombres", "Puesto", "Departamento", "Email", "Extension"])
        for fila in _FILAS_DIRECTORIO[1:]:
            escritor.writerow(fila)


def _generar_csv_vacio() -> None:
    (FIXTURES / "csv_vacio.csv").write_text("", encoding="utf-8")


def _generar_corruptos() -> None:
    (FIXTURES / "corrupto.pdf").write_bytes(b"esto no es un PDF valido, solo texto plano")
    (FIXTURES / "corrupto.xlsx").write_bytes(b"PK\x03\x04esto-no-es-un-zip-valido-de-verdad")
    (FIXTURES / "corrupto.csv").write_bytes(b"\xff\xfe\x00\x01\x02corrupcion-binaria-no-utf8\x80\x81")


def _generar_xls_solo_firma() -> None:
    """`.xls` legado: no hay libreria pura-Python para GENERAR uno real
    sin dependencias adicionales (`xlwt` no soporta Python 3.13). Se deja
    solo la firma OLE -- suficiente para probar la deteccion de
    corrupcion (`xlrd` debe rechazarlo con un error controlado), no la
    extraccion exitosa (cubierta con `xlrd` mockeado en tests unitarios,
    ver 'mejoras futuras' del informe del Modulo 4)."""
    (FIXTURES / "legado_incompleto.xls").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 50)


def generar() -> None:
    _generar_pdf_con_tabla()
    _generar_pdf_con_texto()
    _generar_pdf_escaneado()
    _generar_pdf_vacio()
    _generar_xlsx()
    _generar_xlsx_columnas_desconocidas()
    _generar_xlsx_vacio()
    _generar_csv()
    _generar_csv_vacio()
    _generar_corruptos()
    _generar_xls_solo_firma()


if __name__ == "__main__":
    generar()
    print("Fixtures generadas en", FIXTURES)
