"""Tests para el parser de Excel robusto.

Verifica que el parser maneja correctamente:
- Encabezados en fila 1, 4, etc.
- Encabezados multilínea
- Hoja incorrecta primero
- Filas vacías
- Celdas combinadas
"""

from __future__ import annotations

import openpyxl
import pytest

from govdata.domain.rules.reglas_configurables import ReglasNormalizacion
from govdata.infrastructure.parsers.excel.xlsx_parser import XlsxParser


def _crear_xlsx_con_encabezado_desplazado(ruta: str, fila_encabezado: int = 4) -> None:
    """Crea un archivo XLSX con el encabezado en la fila especificada."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Distributivo"

    # Agregar filas vacías y títulos institucionales antes del encabezado
    for i in range(1, fila_encabezado):
        ws.cell(row=i, column=1, value=f"Título institucional {i}")

    # Encabezado en la fila especificada
    encabezados = ["Apellidos y Nombres", "Cargo", "Dependencia", "Correo", "Teléfono"]
    for col, encabezado in enumerate(encabezados, start=1):
        ws.cell(row=fila_encabezado, column=col, value=encabezado)

    # Datos
    datos = [
        ["PEREZ GOMEZ JUAN CARLOS", "Director Jurídico", "Dirección de Asesoría Jurídica", "jperez@entidad.gob.ec", "02-1234567"],
        ["AYALA SANTI DEXI VIRGINIA", "Procuradora Síndica", "Procuraduría Síndica", "dayala@entidad.gob.ec", "02-7654321"],
        ["RODRIGUEZ MARTINEZ MARIA FERNANDA", "Analista Jurídico", "Coordinación Jurídica", "mrodriguez@entidad.gob.ec", "02-1111111"],
    ]
    for fila, datos_fila in enumerate(datos, start=fila_encabezado + 1):
        for col, valor in enumerate(datos_fila, start=1):
            ws.cell(row=fila, column=col, value=valor)

    wb.save(ruta)


def _crear_xlsx_con_varias_hojas(ruta: str) -> None:
    """Crea un archivo XLSX con varias hojas, donde la primera no es el distributivo."""
    wb = openpyxl.Workbook()

    # Hoja 1: Portada (no es el distributivo)
    ws1 = wb.active
    ws1.title = "Portada"
    ws1.cell(row=1, column=1, value="REPÚBLICA DEL ECUADOR")
    ws1.cell(row=2, column=1, value="Entidad de Prueba")
    ws1.cell(row=3, column=1, value="Directorio y Distributivo del Personal")

    # Hoja 2: Distributivo (es la que nos interesa)
    ws2 = wb.create_sheet("Distributivo")
    encabezados = ["Apellidos y Nombres", "Cargo", "Dependencia"]
    for col, encabezado in enumerate(encabezados, start=1):
        ws2.cell(row=1, column=col, value=encabezado)

    datos = [
        ["PEREZ GOMEZ JUAN CARLOS", "Director Jurídico", "Dirección de Asesoría Jurídica"],
        ["AYALA SANTI DEXI VIRGINIA", "Procuradora Síndica", "Procuraduría Síndica"],
    ]
    for fila, datos_fila in enumerate(datos, start=2):
        for col, valor in enumerate(datos_fila, start=1):
            ws2.cell(row=fila, column=col, value=valor)

    # Hoja 3: Notas (no es el distributivo)
    ws3 = wb.create_sheet("Notas")
    ws3.cell(row=1, column=1, value="Notas al pie")

    wb.save(ruta)


def _crear_xlsx_con_encabezados_multilinea(ruta: str) -> None:
    """Crea un archivo XLSX con encabezados multilínea."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Distributivo"

    # Encabezados multilínea
    ws.cell(row=1, column=1, value="Apellidos y\nNombres")
    ws.cell(row=1, column=2, value="Cargo")
    ws.cell(row=1, column=3, value="Dependencia")

    # Datos
    datos = [
        ["PEREZ GOMEZ JUAN CARLOS", "Director Jurídico", "Dirección de Asesoría Jurídica"],
        ["AYALA SANTI DEXI VIRGINIA", "Procuradora Síndica", "Procuraduría Síndica"],
    ]
    for fila, datos_fila in enumerate(datos, start=2):
        for col, valor in enumerate(datos_fila, start=1):
            ws.cell(row=fila, column=col, value=valor)

    wb.save(ruta)


def _crear_xlsx_con_filas_vacias(ruta: str) -> None:
    """Crea un archivo XLSX con filas vacías entre los datos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Distributivo"

    # Encabezado
    encabezados = ["Apellidos y Nombres", "Cargo", "Dependencia"]
    for col, encabezado in enumerate(encabezados, start=1):
        ws.cell(row=1, column=col, value=encabezado)

    # Datos con filas vacías
    ws.cell(row=2, column=1, value="PEREZ GOMEZ JUAN CARLOS")
    ws.cell(row=2, column=2, value="Director Jurídico")
    ws.cell(row=2, column=3, value="Dirección de Asesoría Jurídica")

    # Fila vacía
    ws.cell(row=3, column=1, value="")
    ws.cell(row=3, column=2, value="")
    ws.cell(row=3, column=3, value="")

    ws.cell(row=4, column=1, value="AYALA SANTI DEXI VIRGINIA")
    ws.cell(row=4, column=2, value="Procuradora Síndica")
    ws.cell(row=4, column=3, value="Procuraduría Síndica")

    wb.save(ruta)


def _crear_xlsx_con_celdas_combinadas(ruta: str) -> None:
    """Crea un archivo XLSX con celdas combinadas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Distributivo"

    # Combinar celdas para el título
    ws.merge_cells("A1:C1")
    ws.cell(row=1, column=1, value="Directorio y Distributivo del Personal")

    # Encabezado
    encabezados = ["Apellidos y Nombres", "Cargo", "Dependencia"]
    for col, encabezado in enumerate(encabezados, start=1):
        ws.cell(row=2, column=col, value=encabezado)

    # Datos
    datos = [
        ["PEREZ GOMEZ JUAN CARLOS", "Director Jurídico", "Dirección de Asesoría Jurídica"],
        ["AYALA SANTI DEXI VIRGINIA", "Procuradora Síndica", "Procuraduría Síndica"],
    ]
    for fila, datos_fila in enumerate(datos, start=3):
        for col, valor in enumerate(datos_fila, start=1):
            ws.cell(row=fila, column=col, value=valor)

    wb.save(ruta)


class TestXlsxParserRobusto:
    """Tests para el parser de Excel robusto."""

    @pytest.fixture
    def reglas(self) -> ReglasNormalizacion:
        return ReglasNormalizacion(
            sinonimos_por_campo={
                "nombres": ["NOMBRES", "NOMBRE", "NOMBRE COMPLETO"],
                "apellidos": ["APELLIDOS", "APELLIDO"],
                "cargo": ["CARGO", "PUESTO", "DENOMINACIÓN"],
                "dependencia": ["DEPENDENCIA", "UNIDAD", "ÁREA"],
                "correo": ["CORREO", "EMAIL", "CORREO ELECTRÓNICO"],
                "telefono": ["TELEFONO", "TELÉFONO", "EXTENSIÓN"],
                "nombre_completo": ["APELLIDOS Y NOMBRES", "NOMBRES Y APELLIDOS", "SERVIDOR", "FUNCIONARIO"],
            }
        )

    def test_encabezado_en_fila_1(self, reglas: ReglasNormalizacion, tmp_path) -> None:
        """Test que el parser detecta correctamente el encabezado en la fila 1."""
        ruta = tmp_path / "test.xlsx"
        _crear_xlsx_con_encabezado_desplazado(str(ruta), fila_encabezado=1)

        parser = XlsxParser(reglas)
        resultado = parser.extraer(str(ruta))

        assert resultado.estado.value == "exito"
        assert resultado.tabla is not None
        assert len(resultado.tabla.filas) == 3

    def test_encabezado_en_fila_4(self, reglas: ReglasNormalizacion, tmp_path) -> None:
        """Test que el parser detecta correctamente el encabezado en la fila 4."""
        ruta = tmp_path / "test.xlsx"
        _crear_xlsx_con_encabezado_desplazado(str(ruta), fila_encabezado=4)

        parser = XlsxParser(reglas)
        resultado = parser.extraer(str(ruta))

        assert resultado.estado.value == "exito"
        assert resultado.tabla is not None
        assert len(resultado.tabla.filas) == 3

    def test_varias_hojas(self, reglas: ReglasNormalizacion, tmp_path) -> None:
        """Test que el parser selecciona correctamente la hoja con el distributivo."""
        ruta = tmp_path / "test.xlsx"
        _crear_xlsx_con_varias_hojas(str(ruta))

        parser = XlsxParser(reglas)
        resultado = parser.extraer(str(ruta))

        assert resultado.estado.value == "exito"
        assert resultado.tabla is not None
        assert len(resultado.tabla.filas) == 2

    def test_encabezados_multilinea(self, reglas: ReglasNormalizacion, tmp_path) -> None:
        """Test que el parser maneja correctamente los encabezados multilínea."""
        ruta = tmp_path / "test.xlsx"
        _crear_xlsx_con_encabezados_multilinea(str(ruta))

        parser = XlsxParser(reglas)
        resultado = parser.extraer(str(ruta))

        assert resultado.estado.value == "exito"
        assert resultado.tabla is not None
        assert len(resultado.tabla.filas) == 2

    def test_filas_vacias(self, reglas: ReglasNormalizacion, tmp_path) -> None:
        """Test que el parser omite correctamente las filas vacías."""
        ruta = tmp_path / "test.xlsx"
        _crear_xlsx_con_filas_vacias(str(ruta))

        parser = XlsxParser(reglas)
        resultado = parser.extraer(str(ruta))

        assert resultado.estado.value == "exito"
        assert resultado.tabla is not None
        assert len(resultado.tabla.filas) == 2  # Solo 2 filas con datos

    def test_celdas_combinadas(self, reglas: ReglasNormalizacion, tmp_path) -> None:
        """Test que el parser maneja correctamente las celdas combinadas."""
        ruta = tmp_path / "test.xlsx"
        _crear_xlsx_con_celdas_combinadas(str(ruta))

        parser = XlsxParser(reglas)
        resultado = parser.extraer(str(ruta))

        assert resultado.estado.value == "exito"
        assert resultado.tabla is not None
        assert len(resultado.tabla.filas) == 2
