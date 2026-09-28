"""Tests: parsers tabulares con documentos "como los reales" (Submodulo 4.1).

Archivos generados en `tmp_path` con datos sinteticos, reproduciendo lo
que se encontro en el portal DPE y en Excel publicos: CSV en
Windows-1252, encabezados desplazados por titulos, celdas combinadas,
encabezados multilinea y la hoja del distributivo que no es la primera.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from govdata.application.dto.resultado_parseo_dto import EstadoParseo
from govdata.config.rule_engine import cargar_rule_engine
from govdata.domain.rules.normalizador_universal import normalizar_tabla
from govdata.infrastructure.parsers.csv.csv_parser import CsvParser, detectar_codificacion
from govdata.infrastructure.parsers.excel.xlsx_parser import XlsxParser

_ENCABEZADO = ["No.", "Apellidos y Nombres", "Puesto Institucional", "Unidad a la que pertenece", "Correo electrónico institucional"]
_FILA = ["1", "PÉREZ GÓMEZ JUAN CARLOS", "ANALISTA JURÍDICO 1", "DIRECCIÓN DE ASESORÍA JURÍDICA", "juan.perez@ejemplo.gob.ec"]


@pytest.fixture(scope="module")
def normalizacion():
    return cargar_rule_engine().normalizacion


class TestCsvCodificacion:
    def test_csv_windows_1252_se_lee_sin_perder_tildes(self, tmp_path: Path, normalizacion) -> None:
        ruta = tmp_path / "cp1252.csv"
        ruta.write_bytes((";".join(_ENCABEZADO) + "\r\n" + ";".join(_FILA) + "\r\n").encode("cp1252"))

        assert detectar_codificacion(str(ruta)) == "cp1252"
        resultado = CsvParser(normalizacion).extraer(str(ruta))

        assert resultado.estado == EstadoParseo.EXITO
        assert resultado.tabla is not None
        assert resultado.tabla.filas[0]["Puesto Institucional"] == "ANALISTA JURÍDICO 1"

    def test_csv_utf8_con_bom(self, tmp_path: Path) -> None:
        ruta = tmp_path / "utf8.csv"
        ruta.write_bytes(("﻿" + ";".join(_ENCABEZADO) + "\n" + ";".join(_FILA) + "\n").encode("utf-8"))

        assert detectar_codificacion(str(ruta)) == "utf-8-sig"
        resultado = CsvParser().extraer(str(ruta))
        assert resultado.tabla is not None
        assert resultado.tabla.encabezados[0] == "No."

    def test_contenido_binario_sigue_siendo_error(self, tmp_path: Path) -> None:
        ruta = tmp_path / "binario.csv"
        ruta.write_bytes(b"\x00\x01\x02PK\x03\x04basura")
        assert CsvParser().extraer(str(ruta)).estado == EstadoParseo.ERROR


class TestCsvEncabezadoDesplazado:
    def test_titulos_arriba_del_encabezado(self, tmp_path: Path, normalizacion) -> None:
        ruta = tmp_path / "titulos.csv"
        lineas = ["GOBIERNO AUTONOMO DESCENTRALIZADO MUNICIPAL EJEMPLO;;;;", ";;;;", ";".join(_ENCABEZADO), ";".join(_FILA)]
        ruta.write_text("\n".join(lineas), encoding="utf-8")

        resultado = CsvParser(normalizacion).extraer(str(ruta))

        assert resultado.tabla is not None
        assert resultado.tabla.encabezados == tuple(_ENCABEZADO)
        assert resultado.tabla.cantidad_filas == 1

    def test_sin_reglas_conserva_la_primera_fila(self, tmp_path: Path) -> None:
        ruta = tmp_path / "titulos.csv"
        ruta.write_text("Titulo;;\nNombre;Cargo;Unidad\nA;B;C\n", encoding="utf-8")
        resultado = CsvParser().extraer(str(ruta))
        assert resultado.tabla is not None
        assert resultado.tabla.encabezados[0] == "Titulo"


def _libro(tmp_path: Path, nombre: str) -> tuple[openpyxl.Workbook, Path]:
    return openpyxl.Workbook(), tmp_path / nombre


class TestXlsxEncabezado:
    def test_encabezado_en_fila_1(self, tmp_path: Path, normalizacion) -> None:
        libro, ruta = _libro(tmp_path, "fila1.xlsx")
        libro.active.append(_ENCABEZADO)
        libro.active.append(_FILA)
        libro.save(ruta)

        resultado = XlsxParser(normalizacion).extraer(str(ruta))
        assert resultado.tabla is not None
        assert resultado.tabla.cantidad_filas == 1

    def test_encabezado_en_fila_4_con_titulo_combinado_y_filas_vacias(self, tmp_path: Path, normalizacion) -> None:
        libro, ruta = _libro(tmp_path, "fila4.xlsx")
        hoja = libro.active
        hoja.append(["MINISTERIO DE EJEMPLO"])
        hoja.merge_cells("A1:E1")
        hoja.append([])
        hoja.append(["Numeral 2: Directorio y distributivo personal de la entidad"])
        hoja.merge_cells("A3:E3")
        hoja.append(_ENCABEZADO)
        hoja.append(_FILA)
        hoja.append([None, None, None, None, None])
        libro.save(ruta)

        resultado = XlsxParser(normalizacion).extraer(str(ruta))

        assert resultado.tabla is not None
        assert resultado.tabla.encabezados == tuple(_ENCABEZADO)
        assert resultado.tabla.cantidad_filas == 1
        filas = normalizar_tabla(resultado.tabla, normalizacion)
        assert filas[0].cargo == "ANALISTA JURÍDICO 1"

    def test_encabezado_multilinea(self, tmp_path: Path, normalizacion) -> None:
        libro, ruta = _libro(tmp_path, "multilinea.xlsx")
        libro.active.append(["Titulo"])
        libro.active.append(["No.", "Apellidos y\nNombres", "Puesto\nInstitucional", "Unidad a la\nque pertenece"])
        libro.active.append(_FILA[:4])
        libro.save(ruta)

        resultado = XlsxParser(normalizacion).extraer(str(ruta))

        assert resultado.tabla is not None
        assert "Puesto Institucional" in resultado.tabla.encabezados

    def test_hoja_incorrecta_primero(self, tmp_path: Path, normalizacion) -> None:
        libro, ruta = _libro(tmp_path, "hojas.xlsx")
        libro.active.title = "Portada"
        libro.active.append(["Informe de transparencia"])
        distributivo = libro.create_sheet("Distributivo")
        distributivo.append(_ENCABEZADO)
        distributivo.append(_FILA)
        libro.save(ruta)

        resultado = XlsxParser(normalizacion).extraer(str(ruta))

        assert resultado.tabla is not None
        assert resultado.tabla.encabezados == tuple(_ENCABEZADO)
        assert resultado.tabla.cantidad_filas == 1

    def test_columnas_duplicadas_no_se_pierden(self, tmp_path: Path, normalizacion) -> None:
        libro, ruta = _libro(tmp_path, "duplicadas.xlsx")
        libro.active.append(["Apellidos y Nombres", "Correo", "Correo"])
        libro.active.append(["PEREZ GOMEZ JUAN CARLOS", "a@ejemplo.gob.ec", "b@ejemplo.com"])
        libro.save(ruta)

        resultado = XlsxParser(normalizacion).extraer(str(ruta))

        assert resultado.tabla is not None
        assert resultado.tabla.filas[0]["Correo"] == "a@ejemplo.gob.ec"
        assert resultado.tabla.filas[0]["Correo (2)"] == "b@ejemplo.com"
